"""Build an Excel dependency report for the Case fields used by the
Automatic Upgrade from Limit Increase feature (NBA-17560).

Usage: python3 scripts/case_field_dependency_report.py [output.xlsx]
"""
import os
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
SRC = os.path.join(ROOT, "force-app", "main", "default")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "Case_Field_Dependency_Report.xlsx")

# (API name, label, data type, confidence, evidence, role in the feature)
FIELDS = [
    ("Is_Initiated_from_Upgrade_Case__c", "Is Initiated from Upgrade Case", "Checkbox", "High",
     "Compared to true in Apex; bound to a checkbox inputField with onchange", "Flag on the Limit Increase case: also create an Upgrade case after approval"),
    ("Target_Upgrade_Credit_Card_Type__c", "Target Upgrade Credit Card Type", "Text (or unrestricted Picklist)", "Medium",
     "Value comes from a lightning:select fed by API product codes", "Card product the customer upgrades to; copied to cc_Requested_Card_Type__c"),
    ("Is_Initiated_from_Limit_Increase_Case__c", "Is Initiated from Limit Increase Case", "Checkbox", "High",
     "Assigned true in CreateUpgradeCases", "Marks the auto-created Upgrade case"),
    ("Selected_Supplementary_Cards__c", "Selected Supplementary Cards", "Long Text Area / Text", "Medium",
     "Holds ';'-joined dynamic values 'id - embossName'", "Supplementary cards carried over to the Upgrade case"),
    ("Un_Selected_Supplementary_Cards__c", "Un-Selected Supplementary Cards", "Long Text Area / Text", "Medium",
     "Holds ';'-joined dynamic values 'id - embossName'", "Supplementary cards excluded from the upgrade"),
    ("cc_Requested_Credit_Limit__c", "Requested Credit Limit", "Currency / Number", "Medium",
     "Numeric input; onchange handleRequestedLimit", "Requested limit on the Limit Increase form"),
    ("sc_Spending_Limit_Amount__c", "Spending Limit Amount", "Currency / Number", "Medium",
     "Displays the minimum limit of the target card", "Shows the new card's minimum credit limit"),
    ("cc_Current_Credit_Limit__c", "Current Credit Limit", "Currency / Number", "Medium",
     "Assigned numeric values (e.g. 10000)", "Set on the Upgrade case from cc_Approved_Credit_Limit__c"),
    ("cc_Approved_Credit_Limit__c", "Approved Credit Limit", "Currency / Number", "Medium",
     "Assigned numeric literals (500, 1000)", "Read from the Limit Increase case"),
    ("cc_Name_on_the_Card__c", "Name on the Card", "Text", "High",
     "Assigned string literals", "Copied to the Upgrade case"),
    ("cc_Credit_Card_PCI_Number__c", "Credit Card PCI Number", "Text", "High",
     "Assigned string literals / PCI variables", "Copied (swapped) to cc_PCI_Id__c on the Upgrade case"),
    ("cc_PCI_Id__c", "PCI Id", "Text", "High",
     "Assigned string literals; bound to masked card number", "Copied (swapped) to cc_Credit_Card_PCI_Number__c on the Upgrade case"),
    ("cc_Approved_Requested_Card_Type__c", "Approved Requested Card Type", "Picklist / Text", "Medium",
     "Compared to product codes (e.g. 'WORLD_CASHBACK_AND_REWARDS_01')", "Copied to Card_Type__c on the Upgrade case"),
    ("Card_Type__c", "Card Type", "Picklist / Text", "Medium",
     "Assigned product codes", "Current card type on the Upgrade case"),
    ("cc_Requested_Card_Type__c", "Requested Card Type", "Picklist / Text", "Medium",
     "Assigned product codes", "Target card type on the Upgrade case"),
    ("Is_Cash_Collateral__c", "Is Cash Collateral", "Checkbox", "High",
     "Compared to true/false", "Queried in CreateUpgradeCases (not used)"),
    ("Sub_Type__c", "Sub Type", "Picklist (likely dependent on Type)", "High",
     "Assigned picklist-style literals ('Credit Limit Increase')", "Set to 'Credit Card Upgrade - No Limit Increase' on the Upgrade case"),
    ("Case_Nature__c", "Case Nature", "Picklist", "High",
     "Assigned picklist-style literals ('Credit Card', 'RC')", "Set to 'Credit Card' on the Upgrade case"),
    ("Region_Flag__c", "Region Flag", "Picklist / Text (possibly formula from Account)", "Low",
     "Compared to 'Bahrain'; same API name exists on Account and User", "Cash collateral path creates upgrade cases only for Bahrain"),
    ("Sub_Status__c", "Sub Status", "Picklist", "High",
     "Assigned picklist-style literals ('Approved', 'In-Progress')", "Set to 'Approved' when the Limit Increase case closes"),
    ("cc_Application_Stages__c", "Application Stages", "Picklist", "High",
     "Assigned picklist-style literals ('Completed', 'Credit Approval')", "Set to 'Completed' on approval"),
    ("cc_CRM_Stage__c", "CRM Stage", "Picklist / Text", "Medium",
     "Assigned 'APPROVED' / 'DECLINED' / mapping values", "Set to 'APPROVED' on approval"),
    ("cc_Activate_Reminders_for_Pending_Cases__c", "Activate Reminders for Pending Cases", "Checkbox (by naming)", "Low",
     "Only commented-out references remain", "Removed from SOQL as part of NBA-17560"),
]
STANDARD = [
    ("AccountId", "Account ID", "Lookup(Account)", "Set on the Upgrade case"),
    ("ParentId", "Parent Case ID", "Lookup(Case)", "Upgrade case's parent = Limit Increase case"),
    ("RecordTypeId", "Record Type ID", "Lookup(RecordType)", "Credit_Card_Upgrade_Request"),
    ("Type", "Type", "Picklist", "Set to 'Credit Card Upgrade'"),
    ("Origin", "Case Origin", "Picklist", "Set to 'CRM'"),
    ("Status", "Status", "Picklist", "Set to 'Closed' when the Limit Increase case is approved"),
]

EXT_TYPES = {
    ".cls": "Apex Class", ".trigger": "Apex Trigger", ".page": "Visualforce Page",
    ".component": "Visualforce Component",
}


def component_of(path):
    rel = os.path.relpath(path, SRC)
    parts = rel.split(os.sep)
    top, name = parts[0], os.path.basename(path)
    if top == "aura":
        return "Aura Component", parts[1]
    if top == "lwc":
        return "LWC", parts[1]
    ext = os.path.splitext(name)[1]
    if ext in EXT_TYPES:
        kind = EXT_TYPES[ext]
        comp = os.path.splitext(name)[0]
        if kind == "Apex Class" and ("test" in comp.lower()):
            kind = "Apex Test Class"
        return kind, comp
    return None, None


def is_test_class(text):
    return re.search(r"@\s*isTest", text, re.I) is not None


def comment_ranges(text):
    """Character ranges covered by // or /* */ comments (approximate, ignores strings)."""
    ranges = []
    for m in re.finditer(r"/\*.*?\*/|//[^\n]*|<!--.*?-->", text, re.S):
        ranges.append((m.start(), m.end()))
    return ranges


def soql_ranges(text):
    """Character ranges of inline [SELECT ...] queries and 'SELECT ...' dynamic query strings."""
    return [(m.start(), m.end()) for m in re.finditer(r"\[\s*SELECT\b.*?\]|'\s*SELECT\b[^']*'", text, re.S | re.I)]


def in_ranges(pos, ranges):
    return any(a <= pos < b for a, b in ranges)


def usage_kind(line, field, kind, pos_in_line):
    before = line[:pos_in_line]
    after = line[pos_in_line + len(field):]
    if kind in ("Aura Component", "LWC", "Visualforce Page", "Visualforce Component"):
        if "inputField" in line or "outputField" in line or "fieldName" in line:
            return "UI field binding"
        if re.search(r"\b(SELECT|FROM|WHERE)\b", line, re.I):
            return "SOQL"
        return "UI / script reference"
    if re.search(r"\bWHERE\b", before, re.I) or re.search(r"\b(AND|OR)\s*$", before.strip(), re.I):
        return "SOQL filter"
    if re.search(r"\bSELECT\b", line, re.I) or re.match(r"^\s*[\w.,\s/*]*$", line) and "," in line:
        return "SOQL select"
    if re.match(r"\s*=[^=]", after):
        return "Write (assignment)"
    if re.match(r"\s*(==|!=|<|>)", after) or re.search(r"(==|!=)\s*$", before):
        return "Read (condition)"
    if re.match(r"\s*:", after):
        return "Write (sObject constructor)"
    return "Read (value)"


def soql_kind(text, pos, ranges):
    for a, b in ranges:
        if a <= pos < b:
            return "SOQL filter" if re.search(r"\bWHERE\b", text[a:pos], re.I) else "SOQL select"
    return None


def scan():
    rows = []
    patterns = {f[0]: re.compile(r"(?<![A-Za-z0-9_])" + re.escape(f[0]) + r"(?![A-Za-z0-9_])") for f in FIELDS}
    for dirpath, _, files in os.walk(SRC):
        for fn in files:
            path = os.path.join(dirpath, fn)
            if fn.endswith("-meta.xml"):
                continue
            kind, comp = component_of(path)
            if not kind:
                continue
            try:
                text = open(path, encoding="utf-8", errors="replace").read()
            except OSError:
                continue
            if kind == "Apex Class" and is_test_class(text):
                kind = "Apex Test Class"
            hits = [f for f, p in patterns.items() if p.search(text)]
            if not hits:
                continue
            cranges = comment_ranges(text)
            qranges = soql_ranges(text)
            line_starts = [0] + [m.end() for m in re.finditer("\n", text)]
            lines = text.split("\n")
            for field in hits:
                for m in patterns[field].finditer(text):
                    pos = m.start()
                    ln = _bisect(line_starts, pos)
                    line = lines[ln]
                    col = pos - line_starts[ln]
                    commented = in_ranges(pos, cranges)
                    prefix = line[max(0, col - 9):col]
                    cross = "Account." in prefix or "__r." in prefix or "Parent." in prefix
                    rows.append({
                        "field": field, "kind": kind, "comp": comp,
                        "file": os.path.relpath(path, ROOT), "line": ln + 1,
                        "usage": "Commented out" if commented else soql_kind(text, pos, qranges) or usage_kind(line, field, kind, col),
                        "active": "No" if commented else "Yes",
                        "cross": "Yes" if cross else "",
                        "snippet": line.strip()[:250],
                    })
    rows.sort(key=lambda r: ([f[0] for f in FIELDS].index(r["field"]), r["kind"], r["comp"], r["file"], r["line"]))
    return rows


def _bisect(starts, pos):
    lo, hi = 0, len(starts) - 1
    while lo < hi:
        mid = (lo + hi + 1) // 2
        if starts[mid] <= pos:
            lo = mid
        else:
            hi = mid - 1
    return lo


# ---------------------------------------------------------------- styling
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
FILL_RED = PatternFill("solid", fgColor="F8CBAD")
FILL_AMBER = PatternFill("solid", fgColor="FFE699")
FILL_GREEN = PatternFill("solid", fgColor="C6E0B4")
FILL_GREY = PatternFill("solid", fgColor="EDEDED")


def write_table(ws, start_row, headers, data, widths, name, wrap_cols=()):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(data, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = BODY, BORDER
            cell.alignment = WRAP if c in wrap_cols else Alignment(vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    end = start_row + max(len(data), 1)
    ref = f"A{start_row}:{get_column_letter(len(headers))}{end}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3 if len(headers) > 4 else 1)
    return end


def build(rows):
    wb = Workbook()
    n = len(rows) + 1
    rng = lambda col: f"'Dependency Detail'!${col}$2:${col}${n}"
    kinds = ["Apex Class", "Apex Test Class", "Apex Trigger", "Aura Component", "LWC",
             "Visualforce Page", "Visualforce Component"]

    # ---------- 1. Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Case Field Dependency Report: Automatic Upgrade from Limit Increase (NBA-17560)"
    ws["A1"].font = TITLE
    info = [
        ("Repository", "UAT-Project-29Sep2026 (force-app/main/default)"),
        ("Scope", "Case fields read or written by the automatic Upgrade-case creation flow "
                  "(CreateUpgradeCases, creditCardLimitApplication, cashCollateralLimitIncreaseApplication, "
                  "initiateManualOnboarding, CreditCardCashCollateral) plus the standard fields it sets"),
        ("Scanned", "Apex classes, Apex triggers, Aura components, LWC, Visualforce pages and components"),
        ("Data types", "Inferred from code usage: field metadata (objects/Case/fields) is not in this repository. "
                       "Confirm in Setup > Object Manager > Case > Fields & Relationships"),
        ("Declarative metadata", "Flows, Process Builder, Workflow, Assignment, Escalation, Sharing and Validation "
                                 "Rules, Layouts and Reports are NOT in this repository; see the "
                                 "'Declarative Coverage' sheet for how to retrieve them"),
        ("Matching", "Whole-word match on the field API name. Region_Flag__c, Sub_Type__c, Sub_Status__c and "
                     "Card_Type__c also exist on other objects, so some hits refer to Account/User fields; "
                     "obvious cross-object references (Account., Parent., __r.) are flagged"),
    ]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = BOLD
        c = ws.cell(row=i, column=2, value=v)
        c.font, c.alignment = BODY, WRAP
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120

    r0 = 11
    ws.cell(row=r0 - 1, column=1, value="Totals (formulas over 'Dependency Detail')").font = BOLD
    totals = [("Custom fields analysed", f"=COUNTA('Field Inventory'!B2:B{len(FIELDS) + 1})"),
              ("Total references", f"=COUNTA({rng('A')})"),
              ("Active references", f"=COUNTIF({rng('G')},\"Yes\")"),
              ("Commented-out references", f"=COUNTIF({rng('G')},\"No\")")]
    for k in kinds:
        totals.append((f"References in {k}", f"=COUNTIF({rng('B')},\"{k}\")"))
    for i, (k, f) in enumerate(totals, r0):
        ws.cell(row=i, column=1, value=k).font = BODY
        c = ws.cell(row=i, column=2, value=f)
        c.font, c.alignment = BOLD, Alignment(horizontal="left")

    # ---------- 2. Field Inventory
    ws = wb.create_sheet("Field Inventory")
    headers = ["#", "API Name", "Field Label (derived)", "Data Type (inferred)", "Type Confidence",
               "Type Evidence", "Role in Upgrade Feature"] + kinds + ["Total Refs", "Active Refs", "Components Using"]
    comps_by_field = defaultdict(set)
    for r in rows:
        comps_by_field[r["field"]].add(r["comp"])
    data = []
    for i, (api, label, dtype, conf, ev, role) in enumerate(FIELDS, 1):
        rr = i + 1
        row = [i, api, label, dtype, conf, ev, role]
        for k in kinds:
            row.append(f"=COUNTIFS({rng('A')},$B{rr},{rng('B')},\"{k}\")")
        row.append(f"=COUNTIF({rng('A')},$B{rr})")
        row.append(f"=COUNTIFS({rng('A')},$B{rr},{rng('G')},\"Yes\")")
        row.append(len(comps_by_field[api]))
        data.append(row)
    widths = [5, 40, 34, 26, 11, 42, 48] + [11] * len(kinds) + [10, 10, 12]
    write_table(ws, 1, headers, data, widths, "FieldInventory", wrap_cols=(4, 6, 7))
    for rr in range(2, len(FIELDS) + 2):
        conf = ws.cell(row=rr, column=5)
        conf.fill = {"High": FILL_GREEN, "Medium": FILL_AMBER, "Low": FILL_RED}[conf.value]
    note_row = len(FIELDS) + 3
    ws.cell(row=note_row, column=2, value="'Components Using' counts distinct components (static value). "
            "All other counts are live formulas over 'Dependency Detail'.").font = Font(name=FONT, size=9, italic=True)

    # ---------- 3. Standard Fields
    ws = wb.create_sheet("Standard Fields")
    data = [[i, a, l, t, "Standard (High)", role] for i, (a, l, t, role) in enumerate(STANDARD, 1)]
    write_table(ws, 1, ["#", "API Name", "Field Label", "Data Type", "Confidence", "Role in Upgrade Feature"],
                data, [5, 18, 20, 22, 16, 60], "StandardFields", wrap_cols=(6,))
    ws.cell(row=len(STANDARD) + 3, column=2,
            value="Standard fields are referenced throughout the org; per-file dependencies are not listed "
                  "because names like Status and Type also match unrelated objects.").font = Font(name=FONT, size=9, italic=True)

    # ---------- 4. Component Usage (one row per field x component)
    ws = wb.create_sheet("Component Usage")
    agg = {}
    for r in rows:
        key = (r["field"], r["kind"], r["comp"], r["file"])
        a = agg.setdefault(key, {"lines": [], "usage": set(), "active": 0, "cross": 0})
        a["lines"].append(r["line"])
        a["usage"].add(r["usage"])
        a["active"] += r["active"] == "Yes"
        a["cross"] += r["cross"] == "Yes"
    data = []
    for (field, kind, comp, f), a in agg.items():
        status = "Active" if a["active"] else "Commented out only"
        data.append([field, kind, comp, f, len(a["lines"]), a["active"], status,
                     ", ".join(sorted(a["usage"])), ", ".join(map(str, a["lines"][:40])) + (" ..." if len(a["lines"]) > 40 else ""),
                     "Yes" if a["cross"] else ""])
    end = write_table(ws, 1, ["API Name", "Component Type", "Component Name", "File", "Refs", "Active Refs",
                              "Status", "Usage Kinds", "Line Numbers", "Has Cross-Object Ref"],
                      data, [38, 17, 40, 70, 7, 8, 18, 40, 40, 12], "ComponentUsage", wrap_cols=(8, 9))
    for rr in range(2, end + 1):
        if ws.cell(row=rr, column=7).value == "Commented out only":
            for c in range(1, 11):
                ws.cell(row=rr, column=c).fill = FILL_GREY

    # ---------- 5. Dependency Matrix (field x component type, distinct components)
    ws = wb.create_sheet("Dependency Matrix")
    by = defaultdict(set)
    for r in rows:
        by[(r["field"], r["kind"])].add(r["comp"])
    data = []
    for api, *_ in FIELDS:
        row = [api]
        for k in kinds:
            comps = sorted(by[(api, k)])
            row.append("; ".join(comps) if len(comps) <= 15 else f"{len(comps)} components (see Component Usage)")
        data.append(row)
    write_table(ws, 1, ["API Name"] + kinds, data, [38] + [45] * len(kinds), "DependencyMatrix",
                wrap_cols=tuple(range(2, len(kinds) + 2)))

    # ---------- 6. Dependency Detail
    ws = wb.create_sheet("Dependency Detail")
    data = [[r["field"], r["kind"], r["comp"], r["file"], r["line"], r["usage"], r["active"], r["cross"], r["snippet"]]
            for r in rows]
    write_table(ws, 1, ["API Name", "Component Type", "Component Name", "File", "Line", "Usage Kind",
                        "Active", "Cross-Object Ref", "Code Snippet"],
                data, [38, 17, 40, 70, 7, 22, 8, 10, 110], "DependencyDetail")

    # ---------- 7. Declarative coverage
    ws = wb.create_sheet("Declarative Coverage")
    decl = [
        ("Flows / Process Builder", "Flow", "Not in repository", "Flow:*"),
        ("Workflow Rules / Field Updates", "Workflow", "Not in repository", "Workflow:Case"),
        ("Assignment Rules", "AssignmentRules", "Not in repository", "AssignmentRules:Case"),
        ("Escalation Rules", "EscalationRules", "Not in repository", "EscalationRules:Case"),
        ("Auto-Response Rules", "AutoResponseRules", "Not in repository", "AutoResponseRules:Case"),
        ("Sharing Rules", "SharingRules", "Not in repository", "SharingRules:Case"),
        ("Validation Rules", "CustomObject / ValidationRule", "Not in repository", "ValidationRule:Case.*"),
        ("Field definitions & data types", "CustomField", "Not in repository", "CustomField:Case.*"),
        ("Formula / Roll-up fields", "CustomField", "Not in repository", "CustomField:Case.*"),
        ("Record Types & picklist values", "RecordType", "Not in repository", "RecordType:Case.*"),
        ("Page Layouts / Lightning Pages", "Layout / FlexiPage", "Not in repository", "Layout:Case-*, FlexiPage:*"),
        ("Field Sets", "FieldSet", "Not in repository", "FieldSet:Case.*"),
        ("Email Templates", "EmailTemplate", "Not in repository", "EmailTemplate:*"),
        ("Reports / Report Types", "Report / ReportType", "Not in repository", "Report:*"),
        ("Profiles / Permission Sets (FLS)", "Profile / PermissionSet", "Not in repository", "PermissionSet:*"),
        ("Apex Classes / Triggers", "ApexClass / ApexTrigger", "Scanned", "-"),
        ("Aura / LWC", "AuraDefinitionBundle / LightningComponentBundle", "Scanned", "-"),
        ("Visualforce Pages / Components", "ApexPage / ApexComponent", "Scanned", "-"),
    ]
    data = [[a, b, c, d] for a, b, c, d in decl]
    end = write_table(ws, 1, ["Dependency Area", "Metadata Type", "Status in This Report", "Retrieve With (sf project retrieve start -m ...)"],
                      data, [34, 40, 22, 60], "DeclarativeCoverage")
    for rr in range(2, end + 1):
        c = ws.cell(row=rr, column=3)
        c.fill = FILL_GREEN if c.value == "Scanned" else FILL_AMBER
    r = end + 2
    ws.cell(row=r, column=1, value="Org-side dependency check (Tooling API, run in Developer Console or sf CLI):").font = BOLD
    ws.cell(row=r + 1, column=1, value=(
        "SELECT MetadataComponentName, MetadataComponentType, RefMetadataComponentName "
        "FROM MetadataComponentDependency WHERE RefMetadataComponentType = 'CustomField' "
        "AND RefMetadataComponentName IN ('Is_Initiated_from_Upgrade_Case', 'Target_Upgrade_Credit_Card_Type', "
        "'Is_Initiated_from_Limit_Increase_Case')")).font = Font(name="Courier New", size=9)
    ws.cell(row=r + 2, column=1, value=(
        "sf data query --use-tooling-api -o <org-alias> -q \"<query above>\"  "
        "(or use 'Where is this used?' on each field in Object Manager)")).font = Font(name="Courier New", size=9)

    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    rows = scan()
    print(build(rows), len(rows), "references")
