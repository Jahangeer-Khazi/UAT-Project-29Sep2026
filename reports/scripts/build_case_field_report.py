#!/usr/bin/env python3
"""Build the Case field dependency Excel report from case_field_dependencies.py output.

Usage: python build_case_field_report.py <deps.json> <out.xlsx> [commit]
"""
import datetime
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

data = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
COMMIT = sys.argv[3] if len(sys.argv) > 3 else ""

CATEGORIES = [
    ("Code", ["Apex Class", "Apex Trigger", "Aura Component", "Lightning Web Component",
              "Visualforce Page", "Visualforce Component"]),
    ("Automation", ["Flow", "Process Builder", "Workflow Rule", "Workflow Field Update",
                    "Workflow Email Alert", "Workflow Outbound Message", "Approval Process",
                    "Assignment Rule", "Auto-Response Rule", "Escalation Rule", "Entitlement Process"]),
    ("Data Integrity & Security", ["Validation Rule", "Validation Rule (Other Object)", "Sharing Rule",
                                   "Formula Field (Case)", "Formula Field (Other Object)",
                                   "Field Dependency (Picklist)", "Lookup Filter",
                                   "Duplicate Rule", "Matching Rule"]),
    ("Configuration Data", ["Custom Metadata Record"]),
    ("UI & Reporting", ["Page Layout", "Compact Layout", "Quick Action", "List View", "Record Type",
                        "Business Process", "Web Link / Button", "Report Type"]),
]
TYPE_CAT = {t: c for c, ts in CATEGORIES for t in ts}
LOGIC_CATS = ["Code", "Automation", "Data Integrity & Security", "Configuration Data"]
ALL_TYPES = [t for _, ts in CATEGORIES for t in ts]

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
CAT_FILLS = {"Code": "DDEBF7", "Automation": "E2EFDA", "Data Integrity & Security": "FCE4D6",
             "Configuration Data": "EDE2F6", "UI & Reporting": "FFF2CC"}
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda **kw: Font(name=FONT, size=kw.pop("size", 10), **kw)


def header_row(ws, row, headers, fill=HEAD_FILL, color="FFFFFF"):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = F(bold=True, color=color)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER


def base_type(f):
    return f["raw_type"] or f["type"].split("(")[0]


wb = Workbook()

# ---------------------------------------------------------------------------
# Field Dependencies (detail) - written first so other sheets can reference it
# ---------------------------------------------------------------------------
fields = {f["api"]: f for f in data["fields"]}
deps = sorted(data["deps"], key=lambda r: (r["field"].lower(), ALL_TYPES.index(r["type"]) if r["type"] in ALL_TYPES else 99,
                                           r["component"].lower(), r["detail"]))
wd = wb.active
wd.title = "Field Dependencies"
dep_headers = ["Field API Name", "Field Label", "Data Type", "Category", "Metadata Type", "Component Name",
               "Where / How Used", "Component Status", "Confidence", "Detection Method", "Line No(s)",
               "File Path", "Evidence (first match)"]
header_row(wd, 1, dep_headers)
for i, r in enumerate(deps, 2):
    f = fields[r["field"]]
    vals = [r["field"], f["label"], f["type"], TYPE_CAT.get(r["type"], "Other"), r["type"], r["component"],
            r["detail"], r["status"], r["confidence"], r["how"], r["lines"], r["path"],
            r["evidence"].replace("&quot;", '"').replace("&apos;", "'").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")]
    for j, v in enumerate(vals, 1):
        c = wd.cell(row=i, column=j, value=v)
        c.font = F()
NDEP = len(deps) + 1
for col, w in zip("ABCDEFGHIJKLM", [34, 30, 22, 22, 24, 40, 34, 13, 12, 34, 14, 60, 80]):
    wd.column_dimensions[col].width = w
wd.freeze_panes = "B2"
tab = Table(displayName="FieldDependencies", ref="A1:M%d" % NDEP)
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
wd.add_table(tab)
green = PatternFill("solid", fgColor="C6EFCE")
amber = PatternFill("solid", fgColor="FFEB9C")
red = PatternFill("solid", fgColor="FFC7CE")
rng = "I2:I%d" % NDEP
wd.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Confirmed"'], fill=green))
wd.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Probable"'], fill=amber))
wd.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Possible"'], fill=red))

DEP_FIELD = "'Field Dependencies'!$A$2:$A$%d" % NDEP
DEP_CAT = "'Field Dependencies'!$D$2:$D$%d" % NDEP
DEP_TYPE = "'Field Dependencies'!$E$2:$E$%d" % NDEP
DEP_COMP = "'Field Dependencies'!$F$2:$F$%d" % NDEP
DEP_CONF = "'Field Dependencies'!$I$2:$I$%d" % NDEP

# Rows are sorted by field, so each field's dependencies sit in one contiguous block.
# Per-field formulas count within that block, which keeps recalculation fast.
blocks = {}
for i, r in enumerate(deps, 2):
    b = blocks.setdefault(r["field"], [i, i])
    b[1] = i

# Only metadata types that actually occur get their own count column
present = {r["type"] for r in deps}
type_cols = [t for t in ALL_TYPES if t in present]

# ---------------------------------------------------------------------------
# Case Fields
# ---------------------------------------------------------------------------
wf = wb.create_sheet("Case Fields", 0)
base_headers = ["#", "Field Label", "Field API Name", "Data Type", "Base Type", "Standard / Custom", "Required",
                "Unique", "External ID", "Formula", "Controlling Field", "Reference To", "History Tracked",
                "Description", "Help Text"]
sum_headers = ["Total References", "Logic References (Code / Automation / Data / Config)",
               "UI & Reporting References", "Usage Status", "Confirmed References"]
headers = base_headers + sum_headers + type_cols
# category banner row
wf.cell(row=1, column=1, value="Case object fields (force-app/main/default/objects/Case/fields)").font = F(bold=True, size=12)
first_type_col = len(base_headers) + len(sum_headers) + 1
col = first_type_col
for cat, ts in CATEGORIES:
    cols = [t for t in ts if t in present]
    if not cols:
        continue
    wf.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(cols) - 1)
    c = wf.cell(row=1, column=col, value=cat)
    c.font = F(bold=True)
    c.fill = PatternFill("solid", fgColor=CAT_FILLS[cat])
    c.alignment = Alignment(horizontal="center")
    col += len(cols)
header_row(wf, 2, headers)
for j, t in enumerate(type_cols, first_type_col):
    wf.cell(row=2, column=j).fill = PatternFill("solid", fgColor=CAT_FILLS[TYPE_CAT[t]])
    wf.cell(row=2, column=j).font = F(bold=True)

flist = sorted(fields.values(), key=lambda f: (f["custom"], f["api"].lower()))
logic_cols = [get_column_letter(first_type_col + i) for i, t in enumerate(type_cols) if TYPE_CAT[t] in LOGIC_CATS]
ui_cols = [get_column_letter(first_type_col + i) for i, t in enumerate(type_cols) if TYPE_CAT[t] == "UI & Reporting"]
last_type_letter = get_column_letter(first_type_col + len(type_cols) - 1)
first_type_letter = get_column_letter(first_type_col)
nb = len(base_headers)
L = lambda k: get_column_letter(nb + k)  # summary column letters
for i, f in enumerate(flist, 3):
    yn = lambda v: "Yes" if v == "true" else ("No" if v == "false" else "")
    vals = [i - 2, f["label"], f["api"], f["type"], base_type(f),
            "Custom" if f["custom"] else "Standard", yn(f["required"]), yn(f["unique"]), yn(f["external_id"]),
            f["formula"], f["controlling_field"], f["reference_to"], yn(f["track_history"]),
            f["description"], f["help"]]
    for j, v in enumerate(vals, 1):
        c = wf.cell(row=i, column=j, value=v)
        c.font = F()
    wf.cell(row=i, column=nb + 1, value="=SUM(%s%d:%s%d)" % (first_type_letter, i, last_type_letter, i))
    wf.cell(row=i, column=nb + 2, value="=" + "+".join("%s%d" % (c, i) for c in logic_cols))
    wf.cell(row=i, column=nb + 3, value="=" + "+".join("%s%d" % (c, i) for c in ui_cols))
    wf.cell(row=i, column=nb + 4, value='=IF(%s%d=0,"Not referenced",IF(%s%d=0,"UI / reporting only","Used in logic"))'
            % (L(1), i, L(2), i))
    b = blocks.get(f["api"])
    if b:
        rf = "'Field Dependencies'!$A$%d:$A$%d" % tuple(b)
        rt = "'Field Dependencies'!$E$%d:$E$%d" % tuple(b)
        rc = "'Field Dependencies'!$I$%d:$I$%d" % tuple(b)
        wf.cell(row=i, column=nb + 5, value='=COUNTIFS(%s,$C%d,%s,"Confirmed")' % (rf, i, rc))
        for j, t in enumerate(type_cols, first_type_col):
            wf.cell(row=i, column=j, value='=COUNTIFS(%s,$C%d,%s,%s$2)' % (rf, i, rt, get_column_letter(j)))
    else:
        for j in [nb + 5] + list(range(first_type_col, first_type_col + len(type_cols))):
            wf.cell(row=i, column=j, value=0)
    for j in range(nb + 1, first_type_col + len(type_cols)):
        wf.cell(row=i, column=j).font = F()
        wf.cell(row=i, column=j).alignment = Alignment(horizontal="center")
NF = len(flist) + 2
widths = [6, 32, 36, 26, 16, 12, 9, 8, 9, 40, 22, 16, 9, 40, 40, 11, 16, 13, 18, 12]
for k, w in enumerate(widths, 1):
    wf.column_dimensions[get_column_letter(k)].width = w
for j in range(first_type_col, first_type_col + len(type_cols)):
    wf.column_dimensions[get_column_letter(j)].width = 12
wf.row_dimensions[2].height = 60
wf.freeze_panes = "D3"
wf.auto_filter.ref = "A2:%s%d" % (last_type_letter, NF)
status_rng = "%s3:%s%d" % (L(4), L(4), NF)
wf.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"Used in logic"'], fill=green))
wf.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"UI / reporting only"'], fill=amber))
wf.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"Not referenced"'], fill=red))
wf.conditional_formatting.add("%s3:%s%d" % (first_type_letter, last_type_letter, NF),
                              CellIsRule(operator="equal", formula=["0"], font=Font(name=FONT, color="BFBFBF")))

# ---------------------------------------------------------------------------
# Components Scanned
# ---------------------------------------------------------------------------
wc = wb.create_sheet("Components Scanned")
header_row(wc, 1, ["Category", "Metadata Type", "Component Name", "Status", "Notes", "File Path",
                   "Case Fields Referenced", "Confirmed References"])
comps = sorted({(c[0], c[1]): c for c in data["components"]}.values(),
               key=lambda c: (ALL_TYPES.index(c[0]) if c[0] in ALL_TYPES else 99, c[1].lower()))
for i, c in enumerate(comps, 2):
    vals = [TYPE_CAT.get(c[0], "Other"), c[0], c[1], c[3], c[4], c[2]]
    for j, v in enumerate(vals, 1):
        wc.cell(row=i, column=j, value=v).font = F()
    wc.cell(row=i, column=7, value="=COUNTIFS(%s,$B%d,%s,$C%d)" % (DEP_TYPE, i, DEP_COMP, i)).font = F()
    wc.cell(row=i, column=8, value='=COUNTIFS(%s,$B%d,%s,$C%d,%s,"Confirmed")' % (DEP_TYPE, i, DEP_COMP, i, DEP_CONF)).font = F()
NC = len(comps) + 1
for col, w in zip("ABCDEFGH", [24, 26, 48, 12, 40, 70, 14, 14]):
    wc.column_dimensions[col].width = w
wc.freeze_panes = "D2"
wc.auto_filter.ref = "A1:H%d" % NC
wc.conditional_formatting.add("G2:G%d" % NC, CellIsRule(operator="equal", formula=["0"], font=Font(name=FONT, color="BFBFBF")))

# ---------------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Summary", 0)
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 18
ws.column_dimensions["D"].width = 18
ws.column_dimensions["E"].width = 18
ws.column_dimensions["F"].width = 60
ws["A1"] = "Case Object - Field Inventory & Dependency Report"
ws["A1"].font = F(bold=True, size=16, color="1F3864")
ws["A2"] = "Repository: UAT-Project-29Sep2026  |  Source: force-app/main/default  |  Commit: %s  |  Generated: %s" % (
    COMMIT, datetime.date.today().isoformat())
ws["A2"].font = F(italic=True, color="595959")

r = 4
ws.cell(row=r, column=1, value="Field inventory").font = F(bold=True, size=12)
r += 1
header_row(ws, r, ["Measure", "Count"], fill=SUB_FILL, color="000000")
CF = "'Case Fields'"
inv = [
    ("Total Case fields in repository", "=COUNTA(%s!$C$3:$C$%d)" % (CF, NF)),
    ("Standard fields", '=COUNTIF(%s!$F$3:$F$%d,"Standard")' % (CF, NF)),
    ("Custom fields", '=COUNTIF(%s!$F$3:$F$%d,"Custom")' % (CF, NF)),
    ("Formula fields", '=COUNTIF(%s!$D$3:$D$%d,"Formula*")' % (CF, NF)),
    ("Fields used in logic (code / automation / data rules / config)", '=COUNTIF(%s!$%s$3:$%s$%d,"Used in logic")' % (CF, L(4), L(4), NF)),
    ("Fields on UI / reports only", '=COUNTIF(%s!$%s$3:$%s$%d,"UI / reporting only")' % (CF, L(4), L(4), NF)),
    ("Fields not referenced anywhere", '=COUNTIF(%s!$%s$3:$%s$%d,"Not referenced")' % (CF, L(4), L(4), NF)),
    ("Total dependency rows", "=COUNTA(%s)" % DEP_FIELD),
]
for label, fml in inv:
    r += 1
    ws.cell(row=r, column=1, value=label).font = F()
    c = ws.cell(row=r, column=2, value=fml)
    c.font = F(bold=True)
    c.number_format = "#,##0"

r += 2
ws.cell(row=r, column=1, value="Fields by data type (base type)").font = F(bold=True, size=12)
r += 1
header_row(ws, r, ["Base Type", "Fields"], fill=SUB_FILL, color="000000")
base_types = sorted({base_type(f) for f in fields.values()},
                    key=lambda t: -sum(1 for f in fields.values() if (base_type(f)) == t))
for t in base_types:
    r += 1
    ws.cell(row=r, column=1, value=t).font = F()
    ws.cell(row=r, column=2, value='=COUNTIF(%s!$E$3:$E$%d,A%d)' % (CF, NF, r)).font = F()

r += 2
ws.cell(row=r, column=1, value="Dependencies by metadata type").font = F(bold=True, size=12)
r += 1
header_row(ws, r, ["Metadata Type", "Category", "Dependency Rows", "Distinct Fields", "Components Using Case Fields", "Notes"],
           fill=SUB_FILL, color="000000")
notes = {
    "Escalation Rule": "No escalationRules metadata in repository - nothing to scan",
    "Record Type": "Picklist value assignments per record type (configuration, not logic)",
    "Business Process": "Each Case business process governs Status values",
    "Report Type": "Case columns exposed in custom report types",
    "Custom Metadata Record": "Field API names held in custom metadata (e.g. Case_Annex JSON, CC_Application_Status_Mapping)",
    "Apex Class": "Includes test classes (flagged 'Test class' in Where / How Used)",
    "Workflow Email Alert": "Email-field recipients (e.g. ContactEmail)",
    "Duplicate Rule": "No Case duplicate rules in repository",
    "Matching Rule": "No Case matching rules in repository",
    "Workflow Outbound Message": "No Case outbound messages in repository",
    "Lookup Filter": "Case lookup filters referencing $Source fields",
    "Validation Rule (Other Object)": "Other objects' rules referencing Case via lookup (e.g. Case__r.Status)",
    "Formula Field (Other Object)": "Other objects' formulas referencing Case via lookup",
    "Web Link / Button": "Case custom buttons/links using {!Case.Field} merge fields",
}
dist_start = r + 1
for cat, ts in CATEGORIES:
    for t in ts:
        r += 1
        ws.cell(row=r, column=1, value=t).font = F()
        ws.cell(row=r, column=2, value=cat).font = F()
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=CAT_FILLS[cat])
        ws.cell(row=r, column=3, value="=COUNTIF(%s,A%d)" % (DEP_TYPE, r)).font = F()
        # distinct fields with at least one reference of this type
        if t in present:
            colL = get_column_letter(first_type_col + type_cols.index(t))
            ws.cell(row=r, column=4, value='=COUNTIF(%s!$%s$3:$%s$%d,">0")' % (CF, colL, colL, NF)).font = F()
        else:
            ws.cell(row=r, column=4, value=0).font = F()
        ws.cell(row=r, column=5, value='=COUNTIFS(\'Components Scanned\'!$B$2:$B$%d,A%d,\'Components Scanned\'!$G$2:$G$%d,">0")' % (NC, r, NC)).font = F()
        ws.cell(row=r, column=6, value=notes.get(t, "")).font = F(color="595959")
        for cc in (3, 4, 5):
            ws.cell(row=r, column=cc).number_format = "#,##0"
r += 1
ws.cell(row=r, column=1, value="Total").font = F(bold=True)
ws.cell(row=r, column=3, value="=SUM(C%d:C%d)" % (dist_start, r - 1)).font = F(bold=True)
ws.cell(row=r, column=3).number_format = "#,##0"
ws.cell(row=r, column=5, value="=SUM(E%d:E%d)" % (dist_start, r - 1)).font = F(bold=True)

r += 2
ws.cell(row=r, column=1, value="Dependency rows by confidence").font = F(bold=True, size=12)
r += 1
header_row(ws, r, ["Confidence", "Rows", "Meaning"], fill=SUB_FILL, color="000000")
ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
for conf, meaning in (
        ("Confirmed", "Reference is unambiguously to the Case field: Case.X / schema import, SOQL FROM Case, a Case-typed "
                      "variable, a flow element or variable on Case, a Case-context formula or rule criteria, a Case layout/list view, etc."),
        ("Probable", "Custom field API name appears in a component that works with Cases, and no other object in the repo "
                     "has a field with that name - but the Case context could not be proven (e.g. dynamic SOQL strings)."),
        ("Possible", "Custom field API name appears, but another object in the repo has a field with the same name or the "
                     "component shows no Case context. Review before relying on it.")):
    r += 1
    ws.cell(row=r, column=1, value=conf).font = F(bold=True)
    ws.cell(row=r, column=2, value="=COUNTIF(%s,A%d)" % (DEP_CONF, r)).font = F()
    ws.cell(row=r, column=2).number_format = "#,##0"
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    c = ws.cell(row=r, column=3, value=meaning)
    c.font = F()
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 42
    ws.cell(row=r, column=1).fill = {"Confirmed": green, "Probable": amber, "Possible": red}[conf]

r += 2
ws.cell(row=r, column=1, value="Scope & method notes").font = F(bold=True, size=12)
method = [
    "Source is the metadata in this repository only - not a live org. Components that exist only in the org are not covered.",
    "Field list = every file in objects/Case/fields (%d). Standard fields with no metadata file (e.g. CaseNumber, IsClosed, RecordTypeId, "
    "CreatedDate) are not listed. Standard field types are filled from the Salesforce Case schema because their files omit <type>." % len(fields),
    "Apex & triggers: comments stripped; matches Case.X, SOQL on Case/Cases (static and string), Case-typed variables (incl. Trigger.new on Case "
    "triggers, maps of Case), new Case(...) and lookups to Case (e.g. Case__r.X). Other custom-field mentions are graded Probable/Possible.",
    "Flows & Process Builders: Get/Create/Update/Delete Records on Case, $Record on Case-triggered flows, Case-typed variables/loops "
    "(e.g. myVariable_current) and Case lookups. Status column shows the flow version status (Active, Draft, Obsolete...).",
    "Aura / LWC / Visualforce: @salesforce/schema/Case.X imports, {!Case.X}, 'Case.X' strings, fieldName / field-name on Case record "
    "forms, plus API-name mentions (graded by confidence). One row per bundle; File Path shows the first matching file.",
    "Declarative: Case validation rules, workflow rules / field updates / email alerts, assignment & auto-response rules (per rule entry), "
    "sharing rules, approval & entitlement processes, page & compact layouts, list views, record types, business processes, quick actions "
    "that create/update Cases, report types, web links, formula fields, dependent picklists, lookup filters and custom metadata records.",
    "Formulas: string literals removed; tokens after a dot belong to related objects (Account.Type is not Case.Type), while X__r.Y counts "
    "as a reference to lookup field X__c and Parent./Account./Owner. count as ParentId/AccountId/OwnerId.",
    "Escalation rules: the repository has no escalationRules folder, so there is nothing to report. Also not in the repository, so not "
    "scanned: Lightning pages (flexipages), email templates, profiles/permission sets (field-level security), field sets, reports & dashboards.",
    "Regenerate: python reports/scripts/case_field_dependencies.py . deps.json && python reports/scripts/build_case_field_report.py deps.json out.xlsx",
]
for m in method:
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value="• " + m)
    c.font = F()
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 40

r += 2
ws.cell(row=r, column=1, value="Sheets").font = F(bold=True, size=12)
for name, desc in (("Case Fields", "One row per field: label, API name, data type, attributes, totals, usage status and a count per metadata type (filterable)."),
                   ("Field Dependencies", "One row per field x component: where and how the field is used, confidence, line numbers, file path and evidence."),
                   ("Components Scanned", "Every component checked, with the number of Case field references it holds (0 = scanned, no Case field use).")):
    r += 1
    ws.cell(row=r, column=1, value=name).font = F(bold=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.cell(row=r, column=2, value=desc).font = F()

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = sheet.title != "Summary"

wb.save(OUT)
print("saved", OUT, "fields", len(fields), "deps", len(deps), "components", len(comps))
